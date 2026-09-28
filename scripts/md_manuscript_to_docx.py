#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Convert manuscript_v1.md to a submission-style .docx with python-docx.
Times New Roman 11pt body (SimSun eastAsia fallback), double-spaced, page numbers,
real bordered tables, embedded figure PNGs after their legends."""
import os, re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

WS = "C:/Users/lenovo/Documents/kimi/tasks/2026-09-26/23-08-42-cb140ac8"
MD = f"{WS}/manuscript/manuscript_v1.md"
OUT = f"{WS}/manuscript/GLP1R青少年自杀风险_投稿稿件_v1.docx"
FIGS = {
    "1": f"{WS}/results/mr_cis/Fig_developmental.png",
    "2": f"{WS}/results/mr_cis/child_mr_forest.png",
    "3": f"{WS}/results/faers_full_forest.png",
}

TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+?\*|`[^`]+?`)")

def add_runs(p, text, base_bold=False):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = p.add_run(part[1:-1]); r.italic = True; r.bold = base_bold
        elif part.startswith("`") and part.endswith("`"):
            p.add_run(part[1:-1])
        else:
            r = p.add_run(part); r.bold = base_bold
    return p

def set_font(style, ascii_name, ea_name, size, bold=None, color=None):
    style.font.name = ascii_name
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), ea_name)

def add_page_number(paragraph):
    run = paragraph.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.text = "PAGE"
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    run._r.append(f1); run._r.append(it); run._r.append(f2)

doc = Document()
for sec in doc.sections:
    sec.left_margin = sec.right_margin = Inches(1)
    sec.top_margin = sec.bottom_margin = Inches(1)

set_font(doc.styles["Normal"], "Times New Roman", "SimSun", 11)
doc.styles["Normal"].paragraph_format.line_spacing = 2.0
doc.styles["Normal"].paragraph_format.space_after = Pt(4)
set_font(doc.styles["Heading 1"], "Arial", "Microsoft YaHei", 13, bold=True, color=RGBColor(0, 0, 0))
set_font(doc.styles["Heading 2"], "Arial", "Microsoft YaHei", 11.5, bold=True, color=RGBColor(0, 0, 0))
doc.styles["Heading 1"].paragraph_format.space_before = Pt(14)
doc.styles["Heading 1"].paragraph_format.line_spacing = 1.15
doc.styles["Heading 2"].paragraph_format.line_spacing = 1.15

footer_p = doc.sections[0].footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_page_number(footer_p)

lines = open(MD, encoding="utf-8").read().splitlines()
i = 0
title_done = False
pagebreak_before_abstract = False
while i < len(lines):
    ln = lines[i].rstrip()
    if not ln.strip():
        i += 1; continue
    if ln.strip() == "---":
        i += 1; continue
    if ln.startswith("# "):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(ln[2:].strip()); r.bold = True; r.font.size = Pt(15)
        p.paragraph_format.line_spacing = 1.3
        title_done = True
        i += 1; continue
    if ln.startswith("## "):
        head = ln[3:].strip()
        if head == "Abstract" and title_done and not pagebreak_before_abstract:
            doc.add_page_break()
            pagebreak_before_abstract = True
        hp = doc.add_heading("", level=1)
        add_runs(hp, head)
        for r in hp.runs: r.bold = True
        i += 1; continue
    if ln.startswith("### "):
        hp = doc.add_heading("", level=2)
        add_runs(hp, ln[4:].strip())
        for r in hp.runs: r.bold = True
        i += 1; continue
    if ln.lstrip().startswith("|"):
        block = []
        while i < len(lines) and lines[i].lstrip().startswith("|"):
            block.append(lines[i].strip()); i += 1
        rows = [[c.strip() for c in row.strip("|").split("|")]
                for row in block if not re.match(r"^\|[\s\-|:]+\|$", row)]
        if rows:
            tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
            tbl.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci, cell in enumerate(row):
                    cp = tbl.rows[ri].cells[ci].paragraphs[0]
                    cp.paragraph_format.line_spacing = 1.15
                    cp.paragraph_format.space_after = Pt(1)
                    add_runs(cp, cell, base_bold=(ri == 0))
                    for r in cp.runs:
                        r.font.size = Pt(9)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        continue
    if ln.lstrip().startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.line_spacing = 1.4
        add_runs(p, ln.lstrip()[2:])
        i += 1; continue
    # regular paragraph
    p = doc.add_paragraph()
    add_runs(p, ln)
    m = re.match(r"^\*\*Figure (\d)\.", ln)
    if m and m.group(1) in FIGS and os.path.exists(FIGS[m.group(1)]):
        ip = doc.add_paragraph()
        ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ip.add_run().add_picture(FIGS[m.group(1)], width=Inches(6.3))
        cp = doc.add_paragraph(f"[Figure {m.group(1)} file: {FIGS[m.group(1)]}]")
        cp.runs[0].font.size = Pt(8); cp.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    i += 1

doc.save(OUT)
print("saved", OUT)
for k, v in FIGS.items():
    print(f"Figure {k}: {'embedded' if os.path.exists(v) else 'MISSING ' + v}")
